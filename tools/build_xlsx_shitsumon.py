# -*- coding: utf-8 -*-
"""インタビュー質問事項を、「なぜ聞くか」だけ空欄にしたエクセルで書き出す。
   実行: python3 tools/build_xlsx_shitsumon.py → pdf/02_インタビュー質問事項_記入用.xlsx
   中身は build_interview.py の PEOPLE をそのまま読む。
   Googleドライブに上げるとスプレッドシートに変換され、そのまま打ち込める。
"""
import os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import build_interview as B

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
JP = "Meiryo"
INK, MUTED, FAINT = "15191C", "5B6266", "8A8F92"
GOLD, LINE, HAIR, BAND = "A8672A", "D9D4CB", "EBE6DD", "F4EFE7"
hair = Side(style="thin", color=HAIR)
inkb = Side(style="medium", color=INK)

WIDTHS = [5, 44, 52, 52, 16]
HEAD = ["№", "質 問", "回 答 例", "な ぜ 聞 く か（ こ こ に 書 く ）", "使 う 動 画"]
CJK = r"[　-〿぀-ヿ㐀-鿿＀-￯]"


def despace(s):
    for _ in range(6):
        s = re.sub("(%s) +(%s)" % (CJK, CJK), r"\1\2", s)
    return re.sub(r" {2,}", " ", s).strip()


def plain(s):
    s = re.sub(r"<br\s*/?>", "\n", str(s or ""))
    s = re.sub(r"<[^>]+>", "", s).replace("　", " ")
    return "\n".join(despace(x) for x in s.split("\n")).strip()


def tags(pick):
    return " ".join(re.findall(r'<span class="tag[^"]*">(.*?)</span>', pick or ""))


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "質問事項"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    r = 1
    ws.cell(row=r, column=1, value="BROOKLYN MUSEUM ／ 向島工房").font = \
        Font(name=JP, size=9, bold=True, color=GOLD)
    r += 1
    c = ws.cell(row=r, column=1, value="インタビュー　質問事項")
    c.font = Font(name=JP, size=18, bold=True, color=INK)
    ws.row_dimensions[r].height = 30
    ws.cell(row=r, column=4, value="「なぜ聞くか」は空けてあります。ここに書き込んでください。").font = \
        Font(name=JP, size=10, bold=True, color=GOLD)
    for i in range(1, 6):
        ws.cell(row=r, column=i).border = Border(bottom=inkb)
    r += 2

    def band(text, size=11, fill=BAND):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cc = ws.cell(row=r, column=1, value=" " + plain(text))
        cc.font = Font(name=JP, size=size, bold=True, color=INK)
        cc.alignment = Alignment(vertical="center", wrap_text=True)
        for i in range(1, 6):
            ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[r].height = 24 if size <= 11 else 30
        r += 1

    def head():
        nonlocal r
        for i, lab in enumerate(HEAD, 1):
            cc = ws.cell(row=r, column=i, value=lab)
            cc.font = Font(name=JP, size=9, color=FAINT if i != 4 else GOLD, bold=(i == 4))
            cc.border = Border(bottom=inkb)
            cc.alignment = Alignment(vertical="bottom")
        ws.row_dimensions[r].height = 20
        r += 1

    # 四人に共通
    band("四 人 に 共 通　──　撮る前に、その場で伝えること", size=13, fill="EDE6DA")
    for a, b in B.COMMON:
        ws.cell(row=r, column=1, value="・").font = Font(name=JP, size=10, color=GOLD)
        ws.cell(row=r, column=2, value=plain(a)).font = Font(name=JP, size=10, bold=True, color=INK)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        cc = ws.cell(row=r, column=3, value=plain(b))
        cc.font = Font(name=JP, size=10, color=MUTED)
        cc.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        for i in range(1, 6):
            ws.cell(row=r, column=i).border = Border(bottom=hair)
        ws.row_dimensions[r].height = max(28, 14 * (len(plain(b)) // 60 + 1))
        r += 1
    r += 1

    for p in B.PEOPLE:
        band("%s　──　%s　／　%s　%s" % (plain(p["name"]), p["role"], p["video"], p["length"]),
             size=13, fill="EDE6DA")
        head()
        i = 0
        for x in p["qs"]:
            if isinstance(x, str):
                band(x, size=10)
                head()
                continue
            i += 1
            q, ex, pick = x
            vals = [i, plain(q), plain(ex), None, tags(pick)]
            for k, v in enumerate(vals, 1):
                cc = ws.cell(row=r, column=k, value=v)
                cc.border = Border(bottom=hair)
                cc.alignment = Alignment(vertical="top", wrap_text=True,
                                         horizontal="center" if k == 1 else "left")
                if k == 1:
                    cc.font = Font(name=JP, size=10, bold=True, color=GOLD)
                elif k == 2:
                    cc.font = Font(name=JP, size=10, bold=True, color=INK)
                elif k == 4:
                    cc.font = Font(name=JP, size=10, color=INK)
                    cc.fill = PatternFill("solid", fgColor="FFFDF7")
                else:
                    cc.font = Font(name=JP, size=10, color=MUTED)
            n = max(len(vals[1]) / 1.7, len(vals[2]) / 2.0)
            ws.row_dimensions[r].height = max(34, 14 * (int(n // 22) + 1))
            r += 1
        r += 1

    return wb


if __name__ == "__main__":
    out = os.path.join(ROOT, "pdf", "02_インタビュー質問事項_記入用.xlsx")
    build().save(out)
    print(out, os.path.getsize(out), "バイト")
