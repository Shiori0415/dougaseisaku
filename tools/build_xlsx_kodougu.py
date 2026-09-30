# -*- coding: utf-8 -*-
"""モデル三人分の小道具・備品を、チェックリストのエクセルで書き出す。
   実行: python3 tools/build_xlsx_kodougu.py → pdf/04_小道具チェックリスト.xlsx
   Googleドライブに上げて「Googleスプレッドシートで開く」と、
   プルダウン（未／済／不要）と色分けが残ったまま開く。
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
JP = "Meiryo"
INK, MUTED, GOLD, LINE, BAND = "15191C", "5B6266", "A8672A", "D9D4CB", "F4EFE7"

PEOPLE = ["難波さん", "上里さん", "北条さん"]
# 各人の欄：○＝必要（未から始める）／任意／要確認／空欄＝不要
ROWS = [
    ("鞄", "鞄（動画用）三パターン", "○", "○", "○", ""),
    ("鞄", "鞄（スチール撮影用・色違いを複数）", "○", "○", "○", "ロケ地の移動中に持って運ぶ"),
    ("鞄", "バッグチャーム", "要確認", "", "", "要否を確認"),
    ("衣装", "衣装 三パターン", "○", "○", "○", ""),
    ("衣装", "靴", "○", "○", "○", ""),
    ("衣装", "アクセサリー", "○", "○", "要確認", "北条さん：絵コンテ（anticrag）で指輪・ネックレスが映る"),
    ("衣装", "サングラス", "要確認", "任意", "要確認",
     "難波さん：絵コンテ（Parisa Wang）でかけている／北条さん：絵コンテ（SELENT ÉTHER）で鞄に入れる"),
    ("衣装", "眼鏡", "", "要確認", "", "上里さん：絵コンテ（chloecleroux）ショット3でかけている"),
    ("小道具", "スマートフォン", "○", "○", "○", "難波さん：操作する場面／上里さん：電話する場面／北条さん：自撮り"),
    ("小道具", "自撮り棒", "", "", "○", ""),
    ("小道具", "飲み物", "", "", "○", ""),
    ("鞄の中身", "ノート・イヤホン・ファイル・手鏡・ポーチ・リップ・メイク道具", "○", "", "○", "鞄に無理なく収まる量"),
    ("鞄の中身", "ノートパソコン・ペンケース", "", "", "○", ""),
    ("鞄の中身", "リップ（動画の中で塗る）または香水（動画の中で吹きかける）", "", "", "○", ""),
    ("雨対策", "傘", "○", "○", "○", ""),
    ("身支度", "鏡（屋外で身なりを確かめる用）", "○", "○", "○", ""),
]
START = {"○": "未", "任意": "任意", "要確認": "要確認", "": "―"}

thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "小道具チェックリスト"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "小道具・備品チェックリスト"
    ws["A1"].font = Font(name=JP, size=14, bold=True, color=INK)
    ws["A2"] = "各人の欄をプルダウンで「済」にする。「―」はその人には要らない物。"
    ws["A2"].font = Font(name=JP, size=9, color=MUTED)

    head = ["分類", "品目"] + PEOPLE + ["備考"]
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(name=JP, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = box

    dv = DataValidation(type="list", formula1='"未,済,不要,任意,要確認,―"', allow_blank=True)
    ws.add_data_validation(dv)

    r = 5
    for i, (kind, item, a, b, c, note) in enumerate(ROWS):
        vals = [kind, item, START[a], START[b], START[c], note]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.font = Font(name=JP, size=10, color=MUTED if v == "―" else INK)
            cell.border = box
            cell.alignment = Alignment(
                horizontal="center" if 3 <= col <= 5 else "left",
                vertical="center", wrap_text=True)
            if i % 2:
                cell.fill = PatternFill("solid", fgColor=BAND)
        for col in (3, 4, 5):
            dv.add(ws.cell(row=r, column=col))
        r += 1

    rng = f"C5:E{r - 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"済"'],
        fill=PatternFill("solid", fgColor="D8EBD3"), font=Font(color="2E6B2E", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"要確認"'],
        fill=PatternFill("solid", fgColor="FBE3C8"), font=Font(color=GOLD, bold=True)))

    for col, w in zip("ABCDEF", [11, 44, 11, 11, 11, 52]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"

    out = os.path.join(ROOT, "pdf", "04_小道具チェックリスト.xlsx")
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
