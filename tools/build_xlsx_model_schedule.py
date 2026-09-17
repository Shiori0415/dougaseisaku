# -*- coding: utf-8 -*-
"""モデル渡しの香盤表（本編と同じ列）を、そのまま読めるエクセルで書き出す。
   実行: python3 tools/build_xlsx_model_schedule.py → pdf/03_香盤表_モデル用_当日.xlsx
   Googleドライブに上げるとGoogleスプレッドシートに変換され、
   列幅・折り返し・帯の色まで残ったまま開く。CSVで上げてはいけない。
   文言は build_model_schedule.py をそのまま読むので、PDF版とずれない。
"""
import os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import build_model_schedule as M

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
JP = "Meiryo"
INK, MUTED, FAINT = "15191C", "5B6266", "8A8F92"
GOLD, LINE, HAIR, BAND = "A8672A", "D9D4CB", "EBE6DD", "F6F2EC"

hair = Side(style="thin", color=HAIR)
line = Side(style="thin", color=LINE)
inkb = Side(style="medium", color=INK)

WIDTHS = [13.5, 30, 20, 58, 7, 40]
HEAD = ["時 刻", "場 所 の 詳 細", "登 場 人 物", "シ ー ン 内 容 詳 細", "尺", "備 考"]


def plain(s):
    s = re.sub(r"<br\s*/?>", "\n", str(s or ""))
    s = re.sub(r"<div[^>]*>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s).replace("　", " ")
    return re.sub(r"\n{2,}", "\n", "\n".join(x.strip() for x in s.split("\n"))).strip()


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "10月13日"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    r = 1

    def span(row, text, size=10, bold=False, color=INK, h=None):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        c = ws.cell(row=row, column=2, value=text)
        c.font = Font(name=JP, size=size, bold=bold, color=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if h:
            ws.row_dimensions[row].height = h
        return c

    ws.cell(row=r, column=1, value="BROOKLYN MUSEUM ／ 向島工房").font = \
        Font(name=JP, size=9, bold=True, color=GOLD)
    ws.row_dimensions[r].height = 18
    r += 1
    c = ws.cell(row=r, column=1, value="撮影スケジュール")
    c.font = Font(name=JP, size=20, bold=True, color=INK)
    c.alignment = Alignment(vertical="center")
    span(r, "%s　／　%s　／　%s" % (M.MODEL, M.DAY, M.SPAN), size=13, bold=True, color=GOLD)
    ws.row_dimensions[r].height = 32
    for i in range(1, 7):
        ws.cell(row=r, column=i).border = Border(bottom=inkb)
    r += 2

    INFO = [("タ イ ト ル", M.TITLE.replace("　", " "), False),
            ("撮 影 日", "%s　9:00集合　　場所 ── %s" % (M.DAY, M.PLACE.replace("　", " ")), True),
            ("時 間", "%s　　出演 ── %s（一日）" % (M.SPAN, M.MODEL), True),
            ("雨 天 予 備 日", "%s　同じ時間・同じ場所" % M.RAIN, True),
            ("緊急時連絡先（宮下）", M.TEL, True)]
    for lab, val, bold in INFO:
        c = ws.cell(row=r, column=1, value=lab)
        c.font = Font(name=JP, size=9, color=FAINT)
        c.alignment = Alignment(vertical="top")
        span(r, val, size=11 if bold else 10, bold=bold,
             color=GOLD if lab.startswith("緊急") else INK)
        for i in range(1, 7):
            ws.cell(row=r, column=i).border = Border(bottom=hair)
        n = len(val)
        ws.row_dimensions[r].height = 22 if n < 96 else 16 * (n // 96 + 1)
        r += 1
    r += 1

    for i, lab in enumerate(HEAD, 1):
        c = ws.cell(row=r, column=i, value=lab)
        c.font = Font(name=JP, size=9, color=FAINT)
        c.alignment = Alignment(vertical="bottom")
        c.border = Border(bottom=inkb)
    ws.row_dimensions[r].height = 20
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    r += 1

    for x in M.ROWS:
        if x[0] == "band":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            c = ws.cell(row=r, column=1, value=" " + plain(x[1]))
            c.font = Font(name=JP, size=11, bold=True, color=INK)
            c.alignment = Alignment(vertical="center")
            for i in range(1, 7):
                ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=BAND)
                ws.cell(row=r, column=i).border = Border(bottom=line)
            ws.row_dimensions[r].height = 26
            r += 1
            continue
        tm, place, who, prop, nm, what, sc, memo = x
        cells = [tm, plain(place), plain(who) + (("\n" + plain(prop)) if prop else ""),
                 plain(nm) + (("\n" + plain(what)) if what else ""), sc, plain(memo)]
        for i, v in enumerate(cells, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = Border(bottom=hair)
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if i == 5 else "left")
            if i == 1:
                c.font = Font(name=JP, size=11, bold=True, color=GOLD)
            elif i == 2:
                c.font = Font(name=JP, size=10, bold=True, color=INK)
            elif i in (5, 6):
                c.font = Font(name=JP, size=10, color=MUTED)
            else:
                c.font = Font(name=JP, size=10, color=INK)
        n = max(len(cells[3]) / 2.4, len(cells[5]) / 1.7, len(cells[2]) / 0.9)
        ws.row_dimensions[r].height = max(34, 14 * (int(n // 24) + 1))
        r += 1

    return wb


if __name__ == "__main__":
    out = os.path.join(ROOT, "pdf", "03_香盤表_モデル用_当日.xlsx")
    build().save(out)
    print(out, os.path.getsize(out), "バイト")
