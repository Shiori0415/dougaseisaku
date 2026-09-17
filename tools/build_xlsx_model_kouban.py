# -*- coding: utf-8 -*-
"""モデル渡し用の香盤表を、そのまま読めるエクセルで書き出す。
   実行: python3 tools/build_xlsx_model_kouban.py → pdf/03_香盤表_モデル用.xlsx
   Googleドライブに上げて「Googleスプレッドシートで開く」と、
   列幅・折り返し・帯の色まで残ったまま開く。
   文言は build_model_kouban.py と同じものを読む。
"""
import os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import build_model_kouban as M

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
JP = "Meiryo"
INK, PROSE, MUTED, FAINT = "15191C", "3D4448", "5B6266", "8A8F92"
GOLD, LINE, HAIR, BAND = "A8672A", "D9D4CB", "EBE6DD", "F4EFE7"

CJK = r"[　-〿぀-ヿ㐀-鿿＀-￯]"
hair = Side(style="thin", color=HAIR)
line = Side(style="thin", color=LINE)
inkb = Side(style="medium", color=INK)

WIDTHS = [13.5, 30, 5, 66, 21]
HEAD = ["時 刻", "場 所", "本", "す る こ と", "衣 装"]


def despace(s):
    for _ in range(6):
        s = re.sub("(%s) +(%s)" % (CJK, CJK), r"\1\2", s)
    return re.sub(r" {2,}", " ", s).strip()


def plain(s):
    s = re.sub(r"<br\s*/?>", "\n", str(s or ""))
    s = re.sub(r"<[^>]+>", "", s).replace("　", " ")
    return "\n".join(despace(x) for x in s.split("\n")).strip()


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "撮影スケジュール"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    r = 1

    def span(row, text, **kw):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        c = ws.cell(row=row, column=2, value=text)
        c.font = Font(name=JP, size=kw.get("size", 10), bold=kw.get("bold", False),
                      color=kw.get("color", INK))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        return c

    # ── 表題 ─────────────────────────────────────────────
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(row=r, column=1, value="BROOKLYN MUSEUM ／ 向島工房")
    c.font = Font(name=JP, size=9, bold=True, color=GOLD)
    ws.row_dimensions[r].height = 18
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(row=r, column=1, value="撮影スケジュール　　難波遥さん")
    c.font = Font(name=JP, size=20, bold=True, color=INK)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 34
    for i in range(1, 6):
        ws.cell(row=r, column=i).border = Border(bottom=inkb)
    r += 2

    # ── 頭の情報 ───────────────────────────────────────────
    INFO = [
        ("撮 影 日", M.DAY + "　" + M.HOURS, True),
        ("雨 天 予 備 日", M.RAIN + "　同じ時間", True),
        ("場 所", "一日、表参道で完結します。　BROOKLYN MUSEUM 表参道店（店休日）と、"
                  "表参道けやき並木。移動はありません。", False),
        ("撮 る 動 画", "① 服装ごとに画が変わる　／　④ 鞄を持って街を歩く　／　"
                        "⑤ コーポレートムービー（店舗の場面だけ）", False),
        ("緊急時連絡先（宮下）", "090-1748-0280　　遅れる・迷った・体調が悪い、いつでもこちらへ。", True),
    ]
    for lab, val, bold in INFO:
        c = ws.cell(row=r, column=1, value=lab)
        c.font = Font(name=JP, size=9, color=FAINT)
        c.alignment = Alignment(vertical="center")
        v = span(r, val, size=11 if bold else 10, bold=bold,
                 color=GOLD if lab.startswith("緊急") else INK)
        for i in range(1, 6):
            ws.cell(row=r, column=i).border = Border(bottom=hair)
        ws.row_dimensions[r].height = 22 if len(val) < 46 else 34
        r += 1
    r += 1

    # ── 当日のこと ─────────────────────────────────────────
    c = ws.cell(row=r, column=1, value="当 日 の こ と")
    c.font = Font(name=JP, size=9, color=FAINT)
    r += 1
    start = r
    for ch in M.CHUI:
        ws.cell(row=r, column=1, value="・").font = Font(name=JP, size=10, color=GOLD)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="top")
        v = span(r, plain(ch), size=10)
        v.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 16 if len(plain(ch)) < 52 else 30
        r += 1
    for i in range(1, 6):
        ws.cell(row=r - 1, column=i).border = Border(bottom=line)
    r += 1

    # ── 本表 ─────────────────────────────────────────────
    def head_row(row):
        for i, lab in enumerate(HEAD, 1):
            c = ws.cell(row=row, column=i, value=lab)
            c.font = Font(name=JP, size=9, color=FAINT)
            c.alignment = Alignment(vertical="bottom")
            c.border = Border(bottom=inkb)
        ws.row_dimensions[row].height = 20

    for x in M.ROWS:
        if x[0] == "band":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            c = ws.cell(row=r, column=1, value="　" + plain(x[1]))
            c.font = Font(name=JP, size=11, bold=True, color=INK)
            c.alignment = Alignment(vertical="center")
            for i in range(1, 6):
                ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=BAND)
                ws.cell(row=r, column=i).border = Border(bottom=Side(style="thin", color=LINE))
            ws.row_dimensions[r].height = 26
            r += 1
            head_row(r)
            r += 1
            continue
        tm, place, no, what, cloth = x
        cl = cloth.replace("</b>\u3000<span class='sm'>", "</b>\n<span>")
        vals = [tm, plain(place), no, plain(what), plain(cl)]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = Border(bottom=hair)
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if i == 3 else "left")
            if i == 1:
                c.font = Font(name=JP, size=10, bold=True, color=GOLD)
            elif i == 2:
                c.font = Font(name=JP, size=10, bold=True, color=INK)
            elif i == 3:
                c.font = Font(name=JP, size=12, bold=True, color=GOLD)
            elif i == 5:
                c.font = Font(name=JP, size=10, color=MUTED)
            else:
                c.font = Font(name=JP, size=10, color=INK)
        n = len(plain(what))
        ws.row_dimensions[r].height = max(30, 16 * (n // 44 + 1))
        r += 1

    return wb


if __name__ == "__main__":
    out = os.path.join(ROOT, "pdf", "03_香盤表_モデル用.xlsx")
    build().save(out)
    print(out, os.path.getsize(out), "バイト")
