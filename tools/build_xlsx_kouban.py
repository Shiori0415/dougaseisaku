# -*- coding: utf-8 -*-
"""香盤表を、シートを分けたエクセルで書き出す。
   実行: python3 tools/build_xlsx_kouban.py → pdf/01_香盤表.xlsx
   Googleドライブに上げて「Googleスプレッドシートで開く」でそのまま見られる。"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import build_shotlist_canvas as L, build_script_canvas as S

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
JP = "Meiryo"
INK, MUTED, GOLD, LINE, BAND, HEADBG = "15191C", "5B6266", "A8672A", "D9D4CB", "F4EFE7", "2E2A24"

thin = Side(style="thin", color=LINE)
BOX = Border(bottom=thin)


def plain(s):
    s = re.sub(r"<br\s*/?>", "\n", str(s or ""))
    return re.sub(r"<[^>]+>", "", s).strip()


def head_row(ws, r, labels, widths):
    for c, (lab, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row=r, column=c, value=lab)
        cell.font = Font(name=JP, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADBG)
        cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[r].height = 22


def put(ws, r, vals, bold=False, band=False, top=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=JP, size=10, bold=bold or (c == 1 and band),
                         color=GOLD if c == 1 else INK)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BOX
        if band:
            cell.fill = PatternFill("solid", fgColor=BAND)


def day_sheet(wb, d):
    ws = wb.create_sheet(plain(d["no"]))
    ws.sheet_view.showGridLines = False
    for i, (k, v) in enumerate([("タイトル", L.day_titles(d)), ("撮影日", d["date2"]),
                                ("場所", d["place"]), ("時間", L.day_span(d)),
                                ("緊急時連絡先", ""), ("注意事項", plain(d["cond"]))], 1):
        a = ws.cell(row=i, column=1, value=k)
        a.font = Font(name=JP, size=9, color=MUTED)
        b = ws.cell(row=i, column=2, value=plain(v))
        b.font = Font(name=JP, size=11 if i <= 2 else 10, bold=i <= 2, color=INK)
        b.alignment = Alignment(vertical="top", wrap_text=True)
    r = 8
    head_row(ws, r, ["時刻", "場所の詳細", "登場人物", "シーン内容詳細", "尺", "備考"],
             [13, 26, 20, 52, 8, 30])
    ws.freeze_panes = "A%d" % (r + 1)
    rows = []
    for tk in d.get("talks", []):
        when, who, place, what, note = tk[:5]
        rows.append((L.start_min(when),
                     [when, plain(place), plain(who) + "\n椅子、レコーダー",
                      plain(tk[6] if len(tk) > 6 else "インタビュー"), "三十分", plain(note)], True))
    for setup in d.get("setups", []):
        when, place, sc, note = setup[:4]
        what = setup[4] if len(setup) > 4 else ""
        who = setup[5] if len(setup) > 5 else ""
        if not sc:
            a, b = [int(x.split(":")[0]) * 60 + int(x.split(":")[1]) for x in when.split("〜")]
            ln = "%d時間%d分" % ((b - a) // 60, (b - a) % 60) if b - a >= 60 else "%d分" % (b - a)
            rows.append((L.start_min(when),
                         [when, plain(place), plain(who) or "―", plain(what) or "―", ln, plain(note)], True))
            continue
        for k, (no, si) in enumerate(sc):
            nm, cuts = L.cuts_of(no, si)
            pl_, _, _, wh = S.SPOT[no][si]
            cam, _ = L.split_cam(S.D[no][si][0])
            w2, prop = L.CASTPROP.get((no, si), ("―", "―"))
            mv, mvday = L.MOVED.get((no, si), ([], ""))
            memo = []
            if k == 0 and note:
                memo.append(plain(note))
            if mv and len(mv) == len(cuts):
                memo.append("%dカット（インタビューの収録で撮る）" % len(cuts))
            else:
                qc, qfrom = L.QUOTE.get((no, si), ([], ""))
                m = "%dカット" % (len(cuts) - len(mv) - len(qc))
                if mv:
                    where = "インタビューの収録で撮る" if mvday == S.SPOT[no][si][2] else "%sに撮る" % mvday
                    m += "（%sは%s）" % ("・".join("%d枚目" % j for j in mv), where)
                if qc:
                    m += "（%sは撮らない ── %s）" % ("・".join("%d枚目" % j for j in qc), qfrom)
                memo.append(m)
            rows.append((L.start_min(wh if wh != "―" else when),
                         [wh, plain(pl_), plain(w2) + "\n" + plain(prop),
                          "%s %s\n%s" % (L.MARU[no], nm, plain(cam)),
                          L.scene_len(no, si), "\n".join(memo)], False))
    rows.sort(key=lambda x: x[0])
    for _, vals, band in rows:
        r += 1
        put(ws, r, vals, band=band)
        ws.row_dimensions[r].height = max(30, 13 * (max(len(str(v).split("\n")) for v in vals) + 1))
    return ws


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "日程"
    ws.sheet_view.showGridLines = False
    t = ws.cell(row=1, column=1, value="BROOKLYN MUSEUM ／ 向島工房　動画八本　香盤表")
    t.font = Font(name=JP, size=14, bold=True, color=INK)
    tot = sum(L.day_counts(d)[0] for d in L.DAYS)
    s2 = ws.cell(row=2, column=1, value="全八本 %dカット ／ 撮影 五日 ／ 日付は10月中で未定" % tot)
    s2.font = Font(name=JP, size=10, color=GOLD)
    head_row(ws, 4, ["日", "場所", "内容", "時間", "カット"], [16, 26, 62, 16, 10])
    ws.freeze_panes = "A5"
    r = 4
    for d in L.DAYS:
        r += 1
        shoot, _ = L.day_counts(d)
        body = "\n".join("%s %s ── %s" % (L.MARU[no], L.pages()[no]["jp"], plain(txt))
                         for no, txt in d["toc"])
        put(ws, r, [d["date"].split("　※")[0], plain(d["place"]), body, L.day_span(d), "%dカット" % shoot])
        ws.row_dimensions[r].height = max(30, 15 * (len(d["toc"]) + 1))
    for d in L.DAYS:
        day_sheet(wb, d)
    cs = wb.create_sheet("出演者と支度")
    cs.sheet_view.showGridLines = False
    head_row(cs, 1, ["本", "撮影日", "当日の支度", "衣装", "そのほか"],
             [8, 20, 36, 36, 44])
    cs.freeze_panes = "A2"
    for i, row in enumerate(L.CAST, 2):
        put(cs, i, [row[0], plain(row[2]), plain(row[3]), plain(row[4]), plain(row[5])])
        cs.row_dimensions[i].height = 58
    out = os.path.join(ROOT, "pdf", "01_香盤表.xlsx")
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
